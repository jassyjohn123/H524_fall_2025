#!/usr/bin/env python3
"""
Extract text content from PowerPoint files for course development.
Uses python-pptx library to read .pptx files.
"""

import os
import sys
from pathlib import Path

from pptx import Presentation

def extract_text_from_pptx(pptx_path):
    """Extract all text content from a PowerPoint file."""
    try:
        prs = Presentation(pptx_path)
        content = []

        for i, slide in enumerate(prs.slides, 1):
            slide_content = [f"\n=== SLIDE {i} ==="]

            # Extract text from all shapes in the slide
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_content.append(shape.text.strip())

                # Handle tables
                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_text = []
                        for cell in row.cells:
                            if cell.text.strip():
                                row_text.append(cell.text.strip())
                        if row_text:
                            slide_content.append(" | ".join(row_text))

            if len(slide_content) > 1:  # Only add if there's content beyond the header
                content.extend(slide_content)

        return "\n".join(content)

    except Exception as e:
        return f"Error reading {pptx_path}: {str(e)}"

def main():
    # Path to the 2024 course materials
    export_path = "/Users/molitorj/ms/OSU/classes/AI_H524/introduction-to-biostatistics-h-524-101-f2024-export"

    # Find all PowerPoint files
    pptx_files = []
    for root, dirs, files in os.walk(export_path):
        for file in files:
            if file.endswith('.pptx'):
                pptx_files.append(os.path.join(root, file))

    print(f"Found {len(pptx_files)} PowerPoint files:")
    for file in pptx_files:
        print(f"  {file}")

    # Extract content from each file
    output_dir = "/Users/molitorj/ms/OSU/classes/AI_H524/H524_fall_2025/extracted_content"
    os.makedirs(output_dir, exist_ok=True)

    for pptx_file in pptx_files:
        print(f"\nProcessing: {os.path.basename(pptx_file)}")
        content = extract_text_from_pptx(pptx_file)

        # Create output filename
        base_name = os.path.splitext(os.path.basename(pptx_file))[0]
        output_file = os.path.join(output_dir, f"{base_name}_content.txt")

        # Write extracted content
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(f"Extracted from: {pptx_file}\n")
            f.write("=" * 50 + "\n\n")
            f.write(content)

        print(f"  Saved to: {output_file}")

    print(f"\nAll content extracted to: {output_dir}")

if __name__ == "__main__":
    main()